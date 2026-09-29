"""
LocalLift — Dedicated Local SEO Intelligence & Audit XLSX Generator

Consumes the exact resolved report dataset from ReportSnapshotService and renders
a complete, multi-sheet Master XLSX Workbook with full row-level data.
"""

import io
from typing import Dict, Any, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class LocalSEOXLSXService:
    @classmethod
    def generate_xlsx(cls, report_data: Dict[str, Any]) -> bytes:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remove default sheet

        # Styles
        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")  # Slate 900
        accent_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")  # Emerald 600
        sub_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")     # Slate 100
        
        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        title_font = Font(name="Arial", size=13, bold=True, color="0F172A")
        sub_title_font = Font(name="Arial", size=10, bold=True, color="475569")
        bold_font = Font(name="Arial", size=9, bold=True, color="0F172A")
        regular_font = Font(name="Arial", size=9, color="1E293B")
        
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        def style_sheet(ws, title_text: str, headers: List[str]):
            ws.views.sheetView[0].showGridLines = True
            ws.append([title_text])
            ws.cell(row=1, column=1).font = title_font
            ws.append([])
            ws.append(headers)
            header_row = 3
            for col_idx, _ in enumerate(headers, 1):
                cell = ws.cell(row=header_row, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.border = thin_border

        def auto_fit_columns(ws):
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.row > 1 and cell.value:
                        val_str = str(cell.value)
                        max_len = max(max_len, len(val_str[:60]))
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        bp = report_data.get("business_profile") or {}
        audit = report_data.get("audit") or {}
        geo = report_data.get("geo_visibility") or {}
        website_audit = report_data.get("website_audit") or {}
        keywords = report_data.get("keywords") or []
        ranking_history = report_data.get("ranking_history") or []
        reviews = (report_data.get("reputation") or {}).get("reviews") or []
        citations = (report_data.get("citations") or {}).get("listings") or []
        competitors = (report_data.get("competitors") or {}).get("items") or []
        schemas = (report_data.get("schema") or {}).get("records") or []
        content_gaps = (report_data.get("content_gaps") or {}).get("opportunities") or []
        provenance = report_data.get("provenance") or {}
        action_plan = report_data.get("action_plan") or {}

        # =====================================================================
        # Sheet 00: Report Metadata
        # =====================================================================
        ws0 = wb.create_sheet(title="00 Report Metadata")
        ws0.views.sheetView[0].showGridLines = True
        ws0.append(["LocalLift — Local SEO Intelligence & Audit Report Metadata"])
        ws0.cell(row=1, column=1).font = title_font
        ws0.append([])
        meta_items = [
            ("Business Name", bp.get("business_name")),
            ("Website / Domain", bp.get("website")),
            ("Report Generated At (UTC)", report_data.get("report_generated_at")),
            ("Authoritative Data As Of (UTC)", report_data.get("data_as_of")),
            ("Overall Local SEO Health Score", f"{report_data.get('health_score')}/100"),
            ("Central Scan Run ID", f"#{ (report_data.get('central_scan') or {}).get('id') or 'N/A' }"),
            ("Central Scan Status", (report_data.get('central_scan') or {}).get('status')),
            ("Standalone Module Overrides", f"{ (report_data.get('overrides_summary') or {}).get('total_overrides', 0) } modules"),
            ("Audit Findings Count", len(audit.get("findings") or [])),
            ("Geo-Grid Points Count", len(geo.get("points") or [])),
            ("Tracked Keywords Count", len(keywords)),
            ("Customer Reviews Count", len(reviews)),
            ("Directory Citations Count", len(citations)),
            ("Competitors Monitored", len(competitors)),
            ("Website Pages Analyzed", len(website_audit.get("pages") or []))
        ]
        ws0.append(["Property / Metric", "Value"])
        ws0.cell(row=3, column=1).fill = header_fill
        ws0.cell(row=3, column=1).font = header_font
        ws0.cell(row=3, column=2).fill = header_fill
        ws0.cell(row=3, column=2).font = header_font
        for k, v in meta_items:
            ws0.append([k, str(v) if v is not None else "N/A"])
            ws0.cell(row=ws0.max_row, column=1).font = bold_font
            ws0.cell(row=ws0.max_row, column=2).font = regular_font
            ws0.cell(row=ws0.max_row, column=1).border = thin_border
            ws0.cell(row=ws0.max_row, column=2).border = thin_border
        auto_fit_columns(ws0)

        # =====================================================================
        # Sheet 01: Executive Summary
        # =====================================================================
        ws1 = wb.create_sheet(title="01 Executive Summary")
        style_sheet(ws1, "Executive Summary & Performance KPIs", ["Metric / Dimension", "Value", "Notes"])
        exec_meta = report_data.get("executive_summary") or {}
        kpis = [
            ("Executive Narrative", exec_meta.get("narrative"), "Overall status"),
            ("Overall Health Score", f"{report_data.get('health_score')}/100", "Composite 20-category framework"),
            ("5x5 Geo-Grid Visibility", f"{geo.get('local_visibility_pct')}%", f"Average rank: {geo.get('average_rank')}"),
            ("Tracked Local Keywords", len(keywords), f"Top 3: {exec_meta.get('metrics', {}).get('top_3_keywords', 0)}"),
            ("Total Customer Reviews", len(reviews), f"Average rating: {exec_meta.get('metrics', {}).get('avg_rating')} ★"),
            ("NAP Citations Consistent", len([c for c in citations if c.get('nap_status') in ('match', 'consistent')]), f"Out of {len(citations)} total listings"),
            ("Technical Site Health", f"{website_audit.get('health_score')}/100", f"{len(website_audit.get('pages', []))} pages analyzed")
        ]
        for m, v, n in kpis:
            ws1.append([m, str(v) if v is not None else "N/A", n])
            r = ws1.max_row
            ws1.cell(row=r, column=1).font = bold_font
            ws1.cell(row=r, column=2).font = regular_font
            ws1.cell(row=r, column=3).font = regular_font
            for c in range(1, 4):
                ws1.cell(row=r, column=c).border = thin_border
        auto_fit_columns(ws1)

        # =====================================================================
        # Sheet 02: Business Profile
        # =====================================================================
        ws2 = wb.create_sheet(title="02 Business Profile")
        style_sheet(ws2, "Canonical Business Profile & NAP Entity", ["Field", "Value"])
        bp_fields = [
            ("Business Name", bp.get("business_name")),
            ("Website URL", bp.get("website")),
            ("Primary Phone", bp.get("primary_phone")),
            ("Primary Address", bp.get("primary_address")),
            ("City", bp.get("city")),
            ("State / Region", bp.get("state")),
            ("Postal Code", bp.get("postal_code")),
            ("Country", bp.get("country")),
            ("Latitude", bp.get("latitude")),
            ("Longitude", bp.get("longitude")),
            ("Primary GBP Category", bp.get("primary_category")),
            ("Additional Categories", ", ".join(bp.get("additional_categories") or [])),
            ("Service Areas", ", ".join(bp.get("service_area") or [])),
            ("Verification Status", bp.get("verification_status")),
            ("Google Place ID", bp.get("place_id")),
            ("Google Maps URL", bp.get("maps_url"))
        ]
        for f, v in bp_fields:
            ws2.append([f, str(v) if v is not None else "N/A"])
            r = ws2.max_row
            ws2.cell(row=r, column=1).font = bold_font
            ws2.cell(row=r, column=2).font = regular_font
            ws2.cell(row=r, column=1).border = thin_border
            ws2.cell(row=r, column=2).border = thin_border
        auto_fit_columns(ws2)

        # =====================================================================
        # Sheet 04: Audit Categories
        # =====================================================================
        ws4 = wb.create_sheet(title="04 Audit Categories")
        style_sheet(ws4, "Local SEO 20-Category Audit Framework Breakdown", [
            "Category Key", "Category Name", "Weight", "Score", "Total Checks", "Passed", "Warnings", "Failed", "Not Verified"
        ])
        for cb in audit.get("category_breakdowns") or []:
            cb_score = cb.get("score")
            if isinstance(cb_score, dict):
                cb_score = cb_score.get("score")
            cb_weight = cb.get("weight", 0)
            if isinstance(cb_weight, dict):
                cb_weight = cb_weight.get("weight", 0)
            try:
                weight_pct = f"{int(float(cb_weight or 0) * 100)}%"
            except (ValueError, TypeError):
                weight_pct = "5%"
            ws4.append([
                cb.get("category_key"),
                cb.get("category_name"),
                weight_pct,
                cb_score if cb_score is not None else "N/A",
                cb.get("total_checks", 0),
                cb.get("passed_checks", 0),
                cb.get("warning_checks", 0),
                cb.get("failed_checks", 0),
                cb.get("not_verified_checks", 0)
            ])
            r = ws4.max_row
            for c in range(1, 10):
                ws4.cell(row=r, column=c).font = regular_font
                ws4.cell(row=r, column=c).border = thin_border
        auto_fit_columns(ws4)

        # =====================================================================
        # Sheet 05: Audit Findings (All Rows)
        # =====================================================================
        ws5 = wb.create_sheet(title="05 Audit Findings")
        style_sheet(ws5, "Complete Local SEO Audit Findings Checkpoints", [
            "Finding ID", "Audit Run ID", "Category Key", "Category Name", "Check Key",
            "Checkpoint Title", "Status", "Severity", "Score Impact", "Evidence",
            "Recommendation", "Source", "Source URL", "Verification Status", "Confidence", "Created At"
        ])
        for f in audit.get("findings") or []:
            ws5.append([
                f.get("id"),
                f.get("audit_run_id"),
                f.get("category_key"),
                f.get("category_name"),
                f.get("check_key"),
                f.get("title"),
                f.get("status"),
                f.get("severity"),
                f.get("score_impact", 0.0),
                f.get("evidence"),
                f.get("recommended_action") or f.get("recommendation"),
                f.get("source"),
                f.get("source_url"),
                f.get("verification_status"),
                f.get("confidence"),
                f.get("created_at")
            ])
            r = ws5.max_row
            for c in range(1, 17):
                ws5.cell(row=r, column=c).font = regular_font
                ws5.cell(row=r, column=c).border = thin_border
        auto_fit_columns(ws5)

        # =====================================================================
        # Sheet 06: Website Pages
        # =====================================================================
        ws6 = wb.create_sheet(title="06 Website Pages")
        style_sheet(ws6, "Technical Website Crawl — Analyzed Pages Graph", [
            "Page ID", "URL", "Status Code", "Title", "Meta Description", "H1", "Word Count",
            "Canonical URL", "Indexable", "Load Time (ms)", "Schema Types", "Images Count",
            "Missing Alt Count", "Internal Links", "External Links"
        ])
        for p in website_audit.get("pages") or []:
            ws6.append([
                p.get("id"),
                p.get("url"),
                p.get("status_code"),
                p.get("title"),
                p.get("meta_description"),
                p.get("h1"),
                p.get("word_count", 0),
                p.get("canonical_url"),
                "YES" if p.get("is_indexable") else "NO",
                p.get("load_time_ms", 0),
                ", ".join(p.get("schema_types") or []),
                p.get("images_count", 0),
                p.get("missing_alt_count", 0),
                p.get("internal_links_count", 0),
                p.get("external_links_count", 0)
            ])
            r = ws6.max_row
            for c in range(1, 16):
                ws6.cell(row=r, column=c).font = regular_font
                ws6.cell(row=r, column=c).border = thin_border
        auto_fit_columns(ws6)

        # =====================================================================
        # Sheet 07: Website Issues
        # =====================================================================
        ws7 = wb.create_sheet(title="07 Website Issues")
        style_sheet(ws7, "Technical & On-Page SEO Issues Detected", [
            "Issue ID", "Category", "Severity", "Title", "Evidence", "Why It Matters",
            "Recommended Solution", "Affected URL", "Status", "Created At"
        ])
        for iss in website_audit.get("issues") or []:
            ws7.append([
                iss.get("id"),
                iss.get("category"),
                iss.get("severity"),
                iss.get("title"),
                iss.get("evidence"),
                iss.get("why_it_matters"),
                iss.get("recommended_solution"),
                iss.get("affected_url"),
                iss.get("status"),
                iss.get("created_at")
            ])
            r = ws7.max_row
            for c in range(1, 11):
                ws7.cell(row=r, column=c).font = regular_font
                ws7.cell(row=r, column=c).border = thin_border
        auto_fit_columns(ws7)

        # =====================================================================
        # Sheet 08: GBP
        # =====================================================================
        ws8 = wb.create_sheet(title="08 GBP")
        style_sheet(ws8, "Google Business Profile Status & Sync", ["Property", "Status / Value"])
        ws8.append(["Connection Status", (provenance.get("gbp") or {}).get("status", "NOT_CONNECTED")])
        ws8.append(["GBP Name", bp.get("business_name")])
        ws8.append(["Place ID", bp.get("place_id") or "Not configured"])
        ws8.append(["Maps URL", bp.get("maps_url") or "Not configured"])
        ws8.append(["Primary Category", bp.get("primary_category") or "Not configured"])
        for r in range(4, 9):
            ws8.cell(row=r, column=1).font = bold_font
            ws8.cell(row=r, column=2).font = regular_font
            ws8.cell(row=r, column=1).border = thin_border
            ws8.cell(row=r, column=2).border = thin_border
        auto_fit_columns(ws8)

        # =====================================================================
        # Sheet 09: Reviews (All Rows)
        # =====================================================================
        ws9 = wb.create_sheet(title="09 Reviews")
        style_sheet(ws9, "Customer Reviews & Reputation Profile", [
            "Review ID", "Author Name", "Rating", "Review Text", "Review Date",
            "Response Text", "Response Status", "Source", "Sentiment", "Created At"
        ])
        for rev in reviews:
            ws9.append([
                rev.get("id"),
                rev.get("author_name"),
                rev.get("rating"),
                rev.get("review_text"),
                rev.get("review_date"),
                rev.get("response_text"),
                rev.get("response_status"),
                rev.get("source"),
                rev.get("sentiment"),
                rev.get("created_at")
            ])
            r = ws9.max_row
            for c in range(1, 11):
                ws9.cell(row=r, column=c).font = regular_font
                ws9.cell(row=r, column=c).border = thin_border
        auto_fit_columns(ws9)

        # =====================================================================
        # Sheet 10: Citations (All Rows)
        # =====================================================================
        ws10 = wb.create_sheet(title="10 Citations")
        style_sheet(ws10, "Directory Citations & External Listings", [
            "Citation ID", "Directory Name", "Domain", "Listing URL", "Domain Authority",
            "Status", "NAP Status", "Verification Status", "Found Name", "Found Address", "Found Phone", "Created At"
        ])
        for cit in citations:
            ws10.append([
                cit.get("id"),
                cit.get("directory_name"),
                cit.get("domain"),
                cit.get("listing_url"),
                cit.get("domain_authority"),
                cit.get("status"),
                cit.get("nap_status"),
                cit.get("verification_status"),
                cit.get("found_name"),
                cit.get("found_address"),
                cit.get("found_phone"),
                cit.get("created_at")
            ])
            r = ws10.max_row
            for c in range(1, 13):
                ws10.cell(row=r, column=c).font = regular_font
                ws10.cell(row=r, column=c).border = thin_border
        auto_fit_columns(ws10)

        # =====================================================================
        # Sheet 11: NAP
        # =====================================================================
        ws11 = wb.create_sheet(title="11 NAP")
        style_sheet(ws11, "NAP Consistency Engine Comparison", ["Metric / Check", "Value"])
        nap_data = report_data.get("nap") or {}
        ws11.append(["NAP Consistency Score", f"{nap_data.get('nap_score')}%"])
        ws11.append(["Total Evaluated Sources", nap_data.get("total_sources_evaluated", len(citations))])
        ws11.append(["Consistent Sources", nap_data.get("consistent_sources_count", 0)])
        ws11.append(["Mismatched Sources", nap_data.get("mismatches_count", 0)])
        for r in range(4, 8):
            ws11.cell(row=r, column=1).font = bold_font
            ws11.cell(row=r, column=2).font = regular_font
            ws11.cell(row=r, column=1).border = thin_border
            ws11.cell(row=r, column=2).border = thin_border
        auto_fit_columns(ws11)

        # =====================================================================
        # Sheet 12: Schema
        # =====================================================================
        ws12 = wb.create_sheet(title="12 Schema")
        style_sheet(ws12, "Schema.org Intelligence & JSON-LD Validation", [
            "Record ID", "Page URL", "Schema Type", "Page Type", "Is Valid",
            "Quality Score", "Detected Types", "Errors", "Warnings", "Last Validated"
        ])
        for s in schemas:
            ws12.append([
                s.get("id"),
                s.get("page_url"),
                s.get("schema_type"),
                s.get("page_type"),
                "TRUE" if s.get("is_valid") else "FALSE",
                s.get("quality_score", 0),
                ", ".join(s.get("detected_types") or []),
                ", ".join(s.get("errors") or []),
                ", ".join(s.get("warnings") or []),
                s.get("last_validated_at")
            ])
            r = ws12.max_row
            for c in range(1, 11):
                ws12.cell(row=r, column=c).font = regular_font
                ws12.cell(row=r, column=c).border = thin_border
        auto_fit_columns(ws12)

        # =====================================================================
        # Sheet 13: Keywords (All Rows)
        # =====================================================================
        ws13 = wb.create_sheet(title="13 Keywords")
        style_sheet(ws13, "Tracked Local Keywords & Current SERP Ranks", [
            "Keyword ID", "Keyword Term", "Search Intent", "Search Volume", "Difficulty",
            "Target Location", "Current Rank", "Previous Rank", "Target Rank",
            "Ranking URL", "SERP Type", "Opportunity Score", "Relevance", "Last Checked"
        ])
        for k in keywords:
            ws13.append([
                k.get("id"),
                k.get("keyword"),
                k.get("search_intent"),
                k.get("search_volume"),
                k.get("difficulty"),
                k.get("target_location"),
                k.get("current_rank"),
                k.get("previous_rank"),
                k.get("target_rank"),
                k.get("ranking_url"),
                k.get("serp_type"),
                k.get("opportunity_score"),
                k.get("business_relevance"),
                k.get("last_checked_at")
            ])
            r = ws13.max_row
            for c in range(1, 15):
                ws13.cell(row=r, column=c).font = regular_font
                ws13.cell(row=r, column=c).border = thin_border
        auto_fit_columns(ws13)

        # =====================================================================
        # Sheet 14: Ranking History
        # =====================================================================
        ws14 = wb.create_sheet(title="14 Ranking History")
        style_sheet(ws14, "Historical Keyword Rank Movement Records", [
            "Keyword ID", "Keyword Term", "Location", "Rank Position", "SERP Type", "Checked At"
        ])
        for h in ranking_history:
            ws14.append([
                h.get("keyword_id"),
                h.get("keyword"),
                h.get("location_name"),
                h.get("rank_position"),
                h.get("serp_type"),
                h.get("checked_at")
            ])
            r = ws14.max_row
            for c in range(1, 7):
                ws14.cell(row=r, column=c).font = regular_font
                ws14.cell(row=r, column=c).border = thin_border
        auto_fit_columns(ws14)

        # =====================================================================
        # Sheet 15: Content Gaps
        # =====================================================================
        ws15 = wb.create_sheet(title="15 Content Gaps")
        style_sheet(ws15, "Content & Suburban Landing Page Opportunities", [
            "Topic / Title", "Target Keyword", "Location", "Search Intent",
            "Opportunity Score", "Recommended Page Type", "Priority"
        ])
        for cg in content_gaps:
            ws15.append([
                cg.get("topic") or cg.get("title"),
                cg.get("target_keyword") or cg.get("keyword"),
                cg.get("location") or cg.get("suburb"),
                cg.get("search_intent", "Commercial"),
                cg.get("opportunity_score"),
                cg.get("recommended_page_type") or cg.get("type"),
                cg.get("priority", "Medium")
            ])
            r = ws15.max_row
            for c in range(1, 8):
                ws15.cell(row=r, column=c).font = regular_font
                ws15.cell(row=r, column=c).border = thin_border
        auto_fit_columns(ws15)

        # =====================================================================
        # Sheet 16: Competitors
        # =====================================================================
        ws16 = wb.create_sheet(title="16 Competitors")
        style_sheet(ws16, "Competitor Benchmarking & Intelligence", [
            "Competitor ID", "Name", "Domain", "Google Place ID", "Rating",
            "Reviews Count", "Citations Count", "Local Visibility Score",
            "Geo-Grid Share %", "Top Keywords Count", "Avg Maps Rank", "Source", "Last Seen"
        ])
        for comp in competitors:
            ws16.append([
                comp.get("id"),
                comp.get("name"),
                comp.get("domain"),
                comp.get("place_id"),
                comp.get("rating"),
                comp.get("reviews_count", 0),
                comp.get("citations_count", 0),
                comp.get("local_visibility_score"),
                comp.get("geo_grid_share_pct"),
                comp.get("top_keywords_count", 0),
                comp.get("avg_maps_rank"),
                comp.get("source"),
                comp.get("last_seen_at")
            ])
            r = ws16.max_row
            for c in range(1, 14):
                ws16.cell(row=r, column=c).font = regular_font
                ws16.cell(row=r, column=c).border = thin_border
        auto_fit_columns(ws16)

        # =====================================================================
        # Sheet 17: GeoGrid Summary
        # =====================================================================
        ws17 = wb.create_sheet(title="17 GeoGrid Summary")
        style_sheet(ws17, "5x5 Geo-Grid Visibility Overview", ["Metric", "Value"])
        ws17.append(["Scan ID", geo.get("scan_id")])
        ws17.append(["Keyword", geo.get("keyword")])
        ws17.append(["Grid Matrix Size", f"{geo.get('grid_size', 5)}x{geo.get('grid_size', 5)} ({geo.get('total_points', 25)} pins)"])
        ws17.append(["Radius (km)", geo.get("radius_km", 10.0)])
        ws17.append(["Local Visibility %", f"{geo.get('local_visibility_pct', 0.0)}%"])
        ws17.append(["Average Rank Position", geo.get("average_rank") or "N/A"])
        ws17.append(["Ranking Found Pins", geo.get("ranking_found_points", 0)])
        ws17.append(["Pins Not Found (20+)", geo.get("not_found_points", 0)])
        ws17.append(["Scan Timestamp", geo.get("scanned_at")])
        for r in range(4, 13):
            ws17.cell(row=r, column=1).font = bold_font
            ws17.cell(row=r, column=2).font = regular_font
            ws17.cell(row=r, column=1).border = thin_border
            ws17.cell(row=r, column=2).border = thin_border
        auto_fit_columns(ws17)

        # =====================================================================
        # Sheet 18: GeoGrid Points (All 25/49 Rows)
        # =====================================================================
        ws18 = wb.create_sheet(title="18 GeoGrid Points")
        style_sheet(ws18, "Geo-Grid Matrix Coordinates & Rank Positions (All Points)", [
            "Point #", "Row", "Col", "Area Name", "Latitude", "Longitude",
            "Keyword", "Rank Position", "Status", "Matched Business", "Matched Place ID",
            "Matched Domain", "Ranking URL", "Distance (km)", "Direction", "Searched At"
        ])
        for pt in geo.get("points") or []:
            ws18.append([
                pt.get("point_number", 0) + 1,
                pt.get("row"),
                pt.get("col"),
                pt.get("area_name"),
                pt.get("latitude"),
                pt.get("longitude"),
                pt.get("keyword"),
                pt.get("rank") if pt.get("rank") is not None else "20+",
                pt.get("status"),
                pt.get("matched_business"),
                pt.get("matched_place_id"),
                pt.get("matched_domain"),
                pt.get("ranking_url"),
                pt.get("distance_km"),
                pt.get("direction"),
                pt.get("searched_at")
            ])
            r = ws18.max_row
            for c in range(1, 17):
                ws18.cell(row=r, column=c).font = regular_font
                ws18.cell(row=r, column=c).border = thin_border
        auto_fit_columns(ws18)

        # =====================================================================
        # Sheet 19: GeoGrid Competitors (All Rows)
        # =====================================================================
        ws19 = wb.create_sheet(title="19 GeoGrid Competitors")
        style_sheet(ws19, "Competitor Rankings Observed Across All Geo-Grid Pins", [
            "Scan ID", "Point ID", "Point #", "Area Name", "Keyword", "Competitor Name",
            "Domain", "Rank Position", "Place ID", "Ranking URL", "Distance (km)"
        ])
        for gc in geo.get("competitors") or []:
            ws19.append([
                gc.get("scan_id"),
                gc.get("point_id"),
                gc.get("point_number", 0) + 1,
                gc.get("area_name"),
                gc.get("keyword"),
                gc.get("competitor_name"),
                gc.get("domain"),
                gc.get("rank"),
                gc.get("place_id"),
                gc.get("ranking_url"),
                gc.get("distance_km")
            ])
            r = ws19.max_row
            for c in range(1, 12):
                ws19.cell(row=r, column=c).font = regular_font
                ws19.cell(row=r, column=c).border = thin_border
        auto_fit_columns(ws19)

        # =====================================================================
        # Sheet 20: Action Plan
        # =====================================================================
        ws20 = wb.create_sheet(title="20 Action Plan")
        style_sheet(ws20, "90-Day Local SEO Action Roadmap", ["Month / Phase", "Focus Strategic Pillar", "Action Items"])
        phases = [
            ("Month 1", action_plan.get("month_1", {}).get("focus", "Foundation & Critical Fixes"), "\n• ".join(action_plan.get("month_1", {}).get("actions", []))),
            ("Month 2", action_plan.get("month_2", {}).get("focus", "On-Page Localization & Suburban Pages"), "\n• ".join(action_plan.get("month_2", {}).get("actions", []))),
            ("Month 3", action_plan.get("month_3", {}).get("focus", "Authority Building & Review Velocity"), "\n• ".join(action_plan.get("month_3", {}).get("actions", [])))
        ]
        for m, f, a in phases:
            ws20.append([m, f, "• " + a if a else "Standardize Local SEO assets."])
            r = ws20.max_row
            ws20.cell(row=r, column=1).font = bold_font
            ws20.cell(row=r, column=2).font = bold_font
            ws20.cell(row=r, column=3).font = regular_font
            ws20.cell(row=r, column=3).alignment = Alignment(wrap_text=True)
            for c in range(1, 4):
                ws20.cell(row=r, column=c).border = thin_border
        auto_fit_columns(ws20)

        # =====================================================================
        # Sheet 21: Provenance
        # =====================================================================
        ws21 = wb.create_sheet(title="21 Provenance")
        style_sheet(ws21, "Report Provenance & Authoritative Module Execution Map", [
            "Module Key", "Authoritative Source Type", "Source Run ID", "Executed At (UTC)", "Execution Status", "Is Standalone Override"
        ])
        for mod_key, prov in provenance.items():
            ws21.append([
                mod_key,
                prov.get("source_type"),
                prov.get("source_run_id"),
                prov.get("executed_at"),
                prov.get("status"),
                "TRUE" if prov.get("is_override") else "FALSE"
            ])
            r = ws21.max_row
            for c in range(1, 7):
                ws21.cell(row=r, column=c).font = regular_font
                ws21.cell(row=r, column=c).border = thin_border
        auto_fit_columns(ws21)

        # Save workbook to memory bytes
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

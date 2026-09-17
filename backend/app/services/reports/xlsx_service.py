import io
import json
from typing import Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

class MasterXLSXService:
    @staticmethod
    def generate_xlsx(canonical_data: Dict[str, Any]) -> bytes:
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Header Styles
        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid") # Indigo
        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        title_font = Font(name="Arial", size=12, bold=True, color="1E293B")
        bold_font = Font(name="Arial", size=9, bold=True, color="0F172A")
        regular_font = Font(name="Arial", size=9, color="334155")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        domain = canonical_data.get("domain", "")
        crawl_id = canonical_data.get("crawl_id") or "N/A"
        timestamp = canonical_data.get("crawl_timestamp") or "N/A"
        analyzed_pages = canonical_data.get("analyzed_pages", 0)
        evaluated_rules = canonical_data.get("evaluated_rules", 0)
        total_checks = canonical_data.get("total_evaluated_checks", 0)

        # =====================================================================
        # Sheet 1: Audit Rules Applied
        # =====================================================================
        ws1 = wb.create_sheet(title="Audit Rules Applied")
        ws1.views.sheetView[0].showGridLines = True

        ws1.append(["Canonical Audit Execution Summary — Master Report"])
        ws1.cell(row=1, column=1).font = title_font

        summary_line = f"Domain: {domain} | Crawl ID: {crawl_id} | Analyzed Pages: {analyzed_pages} | Evaluated Rules: {evaluated_rules} | Total Checks: {total_checks} ({analyzed_pages} pages × {evaluated_rules} rules)"
        ws1.append([summary_line])
        ws1.cell(row=2, column=1).font = bold_font
        ws1.append([])

        headers1 = [
            "Rule ID", "Category", "Rule Name", "Description", "Validation Method",
            "Pages Checked", "Evaluated", "Passed", "Problems", "Status",
            "Requires Integration", "Evidence Source", "Crawl ID", "Crawl Timestamp"
        ]
        ws1.append(headers1)
        for col_num, h_text in enumerate(headers1, 1):
            cell = ws1.cell(row=4, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="left", vertical="center")

        exec_results = canonical_data.get("rule_execution_results", [])
        for r in exec_results:
            row_data = [
                r.get("rule_id", ""),
                r.get("category", ""),
                r.get("rule_name", ""),
                r.get("what_was_checked", ""),
                r.get("validation_method", ""),
                r.get("pages_checked", 0),
                "TRUE" if r.get("evaluated") else "FALSE",
                r.get("passed", 0),
                r.get("problems", 0),
                r.get("status", ""),
                "TRUE" if r.get("requires_integration") else "FALSE",
                r.get("evidence_source", ""),
                crawl_id,
                timestamp
            ]
            ws1.append(row_data)

        # Apply borders & font styling for ws1
        for row in ws1.iter_rows(min_row=4, max_row=ws1.max_row, min_col=1, max_col=14):
            for cell in row:
                cell.border = thin_border
                if cell.row > 4:
                    cell.font = regular_font

        # =====================================================================
        # Sheet 2: Schema Evidence
        # =====================================================================
        ws2 = wb.create_sheet(title="Schema Evidence")
        ws2.views.sheetView[0].showGridLines = True

        headers2 = [
            "Page URL", "Schema Type", "Validation Status", "Completeness",
            "Important Properties", "Missing Properties", "Page Evidence",
            "Raw JSON-LD Available", "Crawl ID"
        ]
        ws2.append(headers2)
        for col_num, h_text in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="left", vertical="center")

        schema_evidence = canonical_data.get("schema_evidence", [])
        for se in schema_evidence:
            props_str = json.dumps(se.get("important_properties", {})) if isinstance(se.get("important_properties"), dict) else str(se.get("important_properties", ""))
            missing_str = ", ".join(se.get("missing_properties", [])) if isinstance(se.get("missing_properties"), list) else str(se.get("missing_properties", ""))
            row_data = [
                se.get("url", ""),
                se.get("schema_type", ""),
                se.get("validation_status", ""),
                se.get("completeness", ""),
                props_str,
                missing_str,
                se.get("page_evidence", ""),
                "TRUE" if se.get("raw_json_ld") else "FALSE",
                crawl_id
            ]
            ws2.append(row_data)

        for row in ws2.iter_rows(min_row=1, max_row=ws2.max_row, min_col=1, max_col=9):
            for cell in row:
                cell.border = thin_border
                if cell.row > 1:
                    cell.font = regular_font

        # =====================================================================
        # Sheet 3: Robots Evidence
        # =====================================================================
        ws3 = wb.create_sheet(title="Robots Evidence")
        ws3.views.sheetView[0].showGridLines = True

        headers3 = [
            "Robots URL", "HTTP Status", "Fetch Status", "User Agent Groups",
            "Allow Rules", "Disallow Rules", "Sitemaps", "Seed URL Result", "Crawl Timestamp"
        ]
        ws3.append(headers3)
        for col_num, h_text in enumerate(headers3, 1):
            cell = ws3.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font

        r_sum = canonical_data.get("robots_summary", {})
        sitemaps_str = ", ".join(r_sum.get("sitemaps", [])) if isinstance(r_sum.get("sitemaps"), list) else str(r_sum.get("sitemaps", ""))
        row_data3 = [
            r_sum.get("robots_url", ""),
            r_sum.get("http_status", ""),
            r_sum.get("fetch_status", ""),
            r_sum.get("user_agent_groups", 0),
            r_sum.get("allow_rules", 0),
            r_sum.get("disallow_rules", 0),
            sitemaps_str,
            r_sum.get("seed_url_result", ""),
            timestamp
        ]
        ws3.append(row_data3)

        for row in ws3.iter_rows(min_row=1, max_row=ws3.max_row, min_col=1, max_col=9):
            for cell in row:
                cell.border = thin_border
                if cell.row > 1:
                    cell.font = regular_font

        # =====================================================================
        # Sheet 4: Health Score
        # =====================================================================
        ws4 = wb.create_sheet(title="Health Score")
        ws4.views.sheetView[0].showGridLines = True

        headers4 = [
            "Metric / Field", "Value / Details"
        ]
        ws4.append(headers4)
        for col_num, h_text in enumerate(headers4, 1):
            cell = ws4.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font

        score_avail = canonical_data.get("score_available", False)
        h_score = canonical_data.get("health_score")
        score_display = str(h_score) if (score_avail and h_score is not None) else "Not Yet Scored"

        score_rows = [
            ["Health Score", score_display],
            ["Score Available", "TRUE" if score_avail else "FALSE"],
            ["Total Evaluated Checks", total_checks],
            ["Passed Checks", canonical_data.get("passed", 0)],
            ["Problem Checks", (canonical_data.get("critical", 0) + canonical_data.get("warnings", 0) + canonical_data.get("opportunities", 0))],
            ["Scoring Formula", canonical_data.get("scoring_formula", "")],
            ["Scoring Weights", json.dumps(canonical_data.get("scoring_weights", {}))]
        ]

        for r_item in score_rows:
            ws4.append(r_item)

        for row in ws4.iter_rows(min_row=1, max_row=ws4.max_row, min_col=1, max_col=2):
            for cell in row:
                cell.border = thin_border
                if cell.row > 1:
                    cell.font = regular_font

        # Auto-adjust column widths
        for ws in [ws1, ws2, ws3, ws4]:
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        buffer = io.BytesIO()
        wb.save(buffer)
        bytes_data = buffer.getvalue()
        buffer.close()
        return bytes_data
